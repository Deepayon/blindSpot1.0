"""Ingestion: CSV, Excel, pytest discovery, and resilience to bad input.

Covers spec §52 items 1-4.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from app.parsers.csv_parser import CsvTestParser
from app.parsers.excel_parser import ExcelTestParser
from app.parsers.jest_parser import JestTestParser
from app.parsers.pytest_parser import PytestParser
from app.parsers.repository_scanner import (
    RepositoryPathError,
    RepositoryScanner,
    validate_repository_path,
)

# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


class TestCsvIngestion:
    def test_parses_canonical_export(self, tmp_path: Path):
        path = tmp_path / "tests.csv"
        path.write_text(
            "Test ID,Test Name,Module,Scenario,Expected Result\n"
            "TC-001,checkout_without_coupon,Checkout,Checkout without coupon,Order succeeds\n"
            "TC-002,checkout_with_discount,Checkout,Checkout with 10% discount,Order succeeds\n",
            encoding="utf-8",
        )
        outcome = CsvTestParser().parse(path)

        assert len(outcome.tests) == 2
        first, second = outcome.tests
        assert first.id == "TC-001"
        assert second.feature == "Checkout"
        assert second.inputs.get("discount") == "10%"
        assert second.expected_behavior == "Order succeeds"

    def test_feature_comes_from_the_export_not_from_guessing_the_text(self, tmp_path: Path):
        """An export states its own taxonomy. Inferring one from row wording
        would reintroduce the domain assumption this release removed."""
        path = tmp_path / "suite.csv"
        path.write_text(
            "ID,Name,Module,Scenario,Expected\n"
            "1,a,Underwriting,Risk assessment for a new policy,Accepted\n"
            "2,b,Claims,Adjudication of an accident report,Settled\n",
            encoding="utf-8",
        )
        features = [t.feature for t in CsvTestParser().parse(path).tests]
        assert features == ["Underwriting", "Claims"]

    def test_without_a_module_column_the_feature_is_not_invented(self, tmp_path: Path):
        path = tmp_path / "checkout_regression.csv"
        path.write_text(
            "ID,Name,Scenario,Expected\n1,a,Some scenario about payments,ok\n", encoding="utf-8"
        )
        test = CsvTestParser().parse(path).tests[0]
        # Derived from the file name, which is real provenance, with the
        # test-type word "regression" dropped as noise. Critically it is not
        # "Payments", which only a guess from the row text would have produced.
        assert test.feature == "Checkout"

    @pytest.mark.parametrize(
        "header",
        [
            "ID,Name,Description,Expected",
            "TestID,Test Case Name,Test Scenario,Expected Behaviour",
            "Key,Title,Steps,Acceptance Criteria",
        ],
    )
    def test_tolerates_column_naming_variations(self, tmp_path: Path, header: str):
        path = tmp_path / "variant.csv"
        path.write_text(f"{header}\nTC-9,login,Login with valid password,Session created\n", encoding="utf-8")
        outcome = CsvTestParser().parse(path)

        assert len(outcome.tests) == 1
        assert outcome.tests[0].id == "TC-9"
        assert outcome.tests[0].name == "login"
        assert outcome.tests[0].expected_behavior == "Session created"

    def test_skips_preamble_rows_above_the_header(self, tmp_path: Path):
        path = tmp_path / "preamble.csv"
        path.write_text(
            "Exported from TestRail,,,\n"
            "2026-01-01,,,\n"
            "Test ID,Test Name,Scenario,Expected Result\n"
            "TC-1,a,Checkout with 20% discount,ok\n",
            encoding="utf-8",
        )
        outcome = CsvTestParser().parse(path)
        assert len(outcome.tests) == 1

    def test_semicolon_delimiter(self, tmp_path: Path):
        path = tmp_path / "semi.csv"
        path.write_text(
            "Test ID;Test Name;Scenario;Expected Result\nTC-1;a;Checkout with tax;ok\n", encoding="utf-8"
        )
        assert len(CsvTestParser().parse(path).tests) == 1

    def test_empty_file_reports_an_issue_and_does_not_raise(self, tmp_path: Path):
        path = tmp_path / "empty.csv"
        path.write_text("", encoding="utf-8")
        outcome = CsvTestParser().parse(path)

        assert outcome.tests == []
        assert any(issue.severity == "ERROR" for issue in outcome.issues)

    def test_unrecognisable_header_is_reported_not_raised(self, tmp_path: Path):
        path = tmp_path / "junk.csv"
        path.write_text("alpha,beta,gamma\n1,2,3\n", encoding="utf-8")
        outcome = CsvTestParser().parse(path)

        assert outcome.tests == []
        assert outcome.files_skipped == 1
        assert "header" in outcome.issues[0].reason.lower()

    def test_explicit_input_column_is_parsed(self, tmp_path: Path):
        path = tmp_path / "inputs.csv"
        path.write_text(
            "Test ID,Test Name,Scenario,Test Data,Expected Result\n"
            "TC-1,a,Checkout,discount=100%; currency=EUR,ok\n",
            encoding="utf-8",
        )
        test = CsvTestParser().parse(path).tests[0]
        assert test.inputs["discount"] == "100%"
        assert test.inputs["currency"] == "EUR"


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------


def _workbook(path: Path, rows: list[list[object]], sheet_title: str = "Checkout") -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


class TestExcelIngestion:
    def test_parses_workbook(self, tmp_path: Path):
        path = _workbook(
            tmp_path / "suite.xlsx",
            [
                ["ID", "Title", "Description", "Expected Behaviour"],
                ["REG-1", "checkout_discount", "Checkout with 20% discount", "Order succeeds"],
                ["REG-2", "checkout_plain", "Checkout without coupon", "Order succeeds"],
            ],
        )
        outcome = ExcelTestParser().parse(path)

        assert len(outcome.tests) == 2
        assert outcome.tests[0].inputs.get("discount") == "20%"
        assert outcome.tests[0].source.endswith("[Checkout]")

    def test_reads_every_sheet(self, tmp_path: Path):
        from openpyxl import Workbook

        path = tmp_path / "multi.xlsx"
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title in ("Checkout", "Payments"):
            sheet = workbook.create_sheet(title)
            sheet.append(["ID", "Title", "Description", "Expected"])
            sheet.append([f"{title}-1", "t", f"{title} scenario", "ok"])
        workbook.save(path)

        outcome = ExcelTestParser().parse(path)
        assert len(outcome.tests) == 2

    def test_title_row_above_header_is_skipped(self, tmp_path: Path):
        path = _workbook(
            tmp_path / "titled.xlsx",
            [
                ["Checkout regression suite"],
                [],
                ["ID", "Title", "Description", "Expected"],
                ["REG-1", "t", "Checkout with tax", "ok"],
            ],
        )
        assert len(ExcelTestParser().parse(path).tests) == 1

    def test_corrupt_workbook_is_reported_not_raised(self, tmp_path: Path):
        path = tmp_path / "corrupt.xlsx"
        path.write_bytes(b"this is definitely not a zip archive")

        outcome = ExcelTestParser().parse(path)
        assert outcome.tests == []
        assert any(issue.severity == "ERROR" for issue in outcome.issues)

    def test_legacy_xls_is_rejected_cleanly(self, tmp_path: Path):
        path = tmp_path / "old.xls"
        path.write_bytes(b"\xd0\xcf\x11\xe0")
        outcome = ExcelTestParser().parse(path)
        assert "xlsx" in outcome.issues[0].reason


# --------------------------------------------------------------------------
# pytest
# --------------------------------------------------------------------------


PYTEST_SOURCE = '''
"""Checkout tests."""
import pytest


def test_checkout_with_a_10_percent_discount():
    """Checkout with a 10% discount coupon.

    Expected: Order total is reduced by 10%
    """
    result = run_scenario(discount="10%")
    assert result.status == "ok"


@pytest.mark.smoke
def test_checkout_without_a_coupon():
    """Checkout without a coupon."""
    result = run_scenario()
    assert result.status == "ok"


@pytest.mark.parametrize("tax", ["0", "5", "8"])
def test_checkout_tax_rates(tax):
    """Checkout across supported tax rates."""
    result = run_scenario(tax=tax)
    assert result.tax == tax


class TestCheckoutEdges:
    def test_checkout_rejects_expired_coupon(self):
        """Checkout with an expired coupon."""
        with pytest.raises(ValueError):
            run_scenario(coupon="expired")


def helper_not_a_test():
    return 1
'''


class TestPytestDiscovery:
    def test_discovers_tests_including_classes(self, tmp_path: Path):
        path = tmp_path / "test_checkout.py"
        path.write_text(PYTEST_SOURCE, encoding="utf-8")
        outcome = PytestParser().parse(path)

        names = {test.name for test in outcome.tests}
        assert names == {
            "test_checkout_with_a_10_percent_discount",
            "test_checkout_without_a_coupon",
            "test_checkout_tax_rates",
            "test_checkout_rejects_expired_coupon",
        }
        assert "helper_not_a_test" not in names

    def test_extracts_inputs_from_call_keywords(self, tmp_path: Path):
        path = tmp_path / "test_checkout.py"
        path.write_text(PYTEST_SOURCE, encoding="utf-8")
        by_name = {test.name: test for test in PytestParser().parse(path).tests}

        assert by_name["test_checkout_with_a_10_percent_discount"].inputs["discount"] == "10%"
        assert by_name["test_checkout_rejects_expired_coupon"].inputs["coupon"] == "expired"

    def test_extracts_parametrize_values(self, tmp_path: Path):
        path = tmp_path / "test_checkout.py"
        path.write_text(PYTEST_SOURCE, encoding="utf-8")
        by_name = {test.name: test for test in PytestParser().parse(path).tests}

        assert by_name["test_checkout_tax_rates"].inputs["tax"] == "0, 5, 8"

    def test_records_provenance_and_markers(self, tmp_path: Path):
        path = tmp_path / "test_checkout.py"
        path.write_text(PYTEST_SOURCE, encoding="utf-8")
        by_name = {test.name: test for test in PytestParser().parse(path).tests}

        plain = by_name["test_checkout_without_a_coupon"]
        assert plain.line_number and plain.line_number > 0
        assert plain.framework == "pytest"
        assert "smoke" in plain.tags
        assert by_name["test_checkout_rejects_expired_coupon"].extra["class"] == "TestCheckoutEdges"

    def test_captures_pytest_raises_as_an_assertion(self, tmp_path: Path):
        path = tmp_path / "test_checkout.py"
        path.write_text(PYTEST_SOURCE, encoding="utf-8")
        by_name = {test.name: test for test in PytestParser().parse(path).tests}
        assert any("raises" in a for a in by_name["test_checkout_rejects_expired_coupon"].assertions)

    def test_malformed_file_is_skipped_not_fatal(self, tmp_path: Path):
        """Spec §43: one bad file must never abort the run."""
        path = tmp_path / "test_broken.py"
        path.write_text("def test_unclosed(:\n    assert True\n", encoding="utf-8")

        outcome = PytestParser().parse(path)
        assert outcome.tests == []
        assert outcome.files_skipped == 1
        assert "Syntax error" in outcome.issues[0].reason

    def test_non_test_file_is_not_supported(self, tmp_path: Path):
        assert PytestParser().supports(tmp_path / "helpers.py") is False
        assert PytestParser().supports(tmp_path / "test_thing.py") is True
        assert PytestParser().supports(tmp_path / "thing_test.py") is True


class TestJestDiscovery:
    def test_discovers_it_and_describe(self, tmp_path: Path):
        path = tmp_path / "checkout.test.ts"
        path.write_text(
            """
describe('Checkout', () => {
  it('applies a 10% discount coupon', () => {
    expect(order.total).toBe(90);
  });
  test('rejects an expired coupon', () => {
    expect(result.error).toBe('expired');
  });
});
""",
            encoding="utf-8",
        )
        outcome = JestTestParser().parse(path)

        assert len(outcome.tests) == 2
        assert outcome.tests[0].extra["suite"] == "Checkout"
        assert outcome.tests[0].inputs.get("discount") == "10%"


# --------------------------------------------------------------------------
# Repository scanning and its security posture (spec §40)
# --------------------------------------------------------------------------


class TestRepositoryScanning:
    def test_scans_and_reports_skipped_files(self, tmp_path: Path):
        tests_dir = tmp_path / "tests"
        tests_dir.mkdir()
        (tests_dir / "test_checkout.py").write_text(PYTEST_SOURCE, encoding="utf-8")
        (tests_dir / "test_broken.py").write_text("def test_x(:\n  pass\n", encoding="utf-8")
        (tmp_path / "README.md").write_text("not a test", encoding="utf-8")

        outcome = RepositoryScanner().scan(tmp_path)

        assert len(outcome.tests) == 4
        assert outcome.files_scanned == 1
        assert outcome.files_skipped == 1
        assert outcome.tests[0].source == "tests/test_checkout.py"

    def test_ignored_directories_are_pruned(self, tmp_path: Path):
        for ignored in ("node_modules", ".venv", "__pycache__"):
            directory = tmp_path / ignored
            directory.mkdir()
            (directory / "test_vendored.py").write_text("def test_v():\n    assert True\n", encoding="utf-8")
        (tmp_path / "test_real.py").write_text("def test_r():\n    assert True\n", encoding="utf-8")

        outcome = RepositoryScanner().scan(tmp_path)
        assert len(outcome.tests) == 1
        assert outcome.tests[0].name == "test_r"

    def test_empty_repository_yields_nothing_without_error(self, tmp_path: Path):
        outcome = RepositoryScanner().scan(tmp_path)
        assert outcome.tests == []

    def test_rejects_remote_urls(self):
        with pytest.raises(RepositoryPathError, match="local directories only"):
            validate_repository_path("https://github.com/example/repo.git")

    def test_rejects_unc_paths(self):
        with pytest.raises(RepositoryPathError, match="Network"):
            validate_repository_path(r"\\server\share\repo")

    def test_rejects_missing_path(self, tmp_path: Path):
        with pytest.raises(RepositoryPathError):
            validate_repository_path(str(tmp_path / "does-not-exist"))

    def test_rejects_a_file(self, tmp_path: Path):
        target = tmp_path / "a.py"
        target.write_text("x = 1", encoding="utf-8")
        with pytest.raises(RepositoryPathError, match="is a file"):
            validate_repository_path(str(target))

    def test_rejects_empty_path(self):
        with pytest.raises(RepositoryPathError):
            validate_repository_path("   ")

    def test_oversized_file_is_skipped(self, tmp_path: Path, monkeypatch):
        from app.config.settings import get_settings

        settings = get_settings()
        monkeypatch.setattr(settings, "max_file_size_bytes", 50)
        (tmp_path / "test_big.py").write_text(PYTEST_SOURCE, encoding="utf-8")

        outcome = RepositoryScanner(settings=settings).scan(tmp_path)
        assert outcome.tests == []
        assert "exceeds" in outcome.issues[0].reason
