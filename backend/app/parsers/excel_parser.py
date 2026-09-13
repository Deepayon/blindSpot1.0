"""Excel (.xlsx/.xlsm) test ingestion via openpyxl.

Every worksheet is scanned independently, so a workbook with one sheet per
feature works without configuration. Sheets without a recognisable header row
are reported and skipped rather than failing the workbook.
"""
from __future__ import annotations

from pathlib import Path

from ..config.logging_conf import get_logger
from .base import ParseOutcome, TestParser
from .tabular import build_test_from_row, has_usable_headers, map_headers

log = get_logger(__name__)

_MAX_HEADER_SCAN = 10
_MAX_ROWS_PER_SHEET = 100_000


class ExcelTestParser(TestParser):
    framework = "excel"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}

    def parse(self, path: Path) -> ParseOutcome:
        outcome = ParseOutcome()

        if path.suffix.lower() == ".xls":
            outcome.warn(
                str(path),
                "Legacy .xls is not supported by openpyxl — re-save as .xlsx.",
                "ERROR",
            )
            outcome.files_skipped = 1
            return outcome

        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover - openpyxl is a declared dependency
            outcome.warn(str(path), "openpyxl is not installed.", "ERROR")
            return outcome

        try:
            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        except Exception as exc:
            outcome.warn(str(path), f"Could not open workbook (corrupt or unsupported): {exc}", "ERROR")
            outcome.files_skipped = 1
            return outcome

        outcome.files_scanned = 1
        try:
            for sheet in workbook.worksheets:
                self._parse_sheet(sheet, path, outcome)
        finally:
            workbook.close()

        if not outcome.tests:
            outcome.warn(str(path), "No test rows found in any worksheet.")

        log.info(
            "excel parsed",
            extra={
                "event": "parse.excel",
                "file": path.name,
                "sheets": len(workbook.sheetnames),
                "tests": len(outcome.tests),
            },
        )
        return outcome

    def _parse_sheet(self, sheet, path: Path, outcome: ParseOutcome) -> None:
        location = f"{path.name}[{sheet.title}]"
        try:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                rows.append(list(row))
                if index >= _MAX_ROWS_PER_SHEET:
                    outcome.warn(location, f"Sheet truncated at {_MAX_ROWS_PER_SHEET} rows.")
                    break
        except Exception as exc:
            outcome.warn(location, f"Sheet could not be read: {exc}", "ERROR")
            return

        if not rows:
            outcome.warn(location, "Sheet is empty.")
            return

        header_index = None
        for index, row in enumerate(rows[:_MAX_HEADER_SCAN]):
            if any(str(cell).strip() for cell in row if cell is not None) and has_usable_headers(row):
                header_index = index
                break

        if header_index is None:
            outcome.warn(location, "No recognisable header row; sheet skipped.")
            return

        mapping = map_headers(rows[header_index])
        id_prefix = f"{path.stem.upper()[:8]}-{sheet.title.upper()[:6]}".replace(" ", "")

        for offset, row in enumerate(rows[header_index + 1 :], start=1):
            if not any(str(cell).strip() for cell in row if cell is not None):
                continue
            try:
                test = build_test_from_row(
                    row,
                    mapping,
                    row_number=header_index + 1 + offset,
                    source=location,
                    framework=self.framework,
                    id_prefix=id_prefix,
                )
            except Exception as exc:
                outcome.warn(f"{location}:row {header_index + 1 + offset}", f"Row skipped: {exc}")
                continue
            if test is not None:
                outcome.tests.append(test)
